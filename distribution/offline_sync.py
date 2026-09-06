"""
Offline sync utilities for S.A.R.A. distribution backup and reconciliation.

This module handles:
1. Beneficiary list export (schedule-specific) for offline backup
2. Full municipal-wide sync export (households, programs, claims)
3. Import on offline instance
4. Claim reconciliation export/import
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.utils import timezone
from django.db.models import Q
from decimal import Decimal
import zipfile
import json
from io import BytesIO

from .models import AidSchedule, GeneratedBeneficiaryList, GeneratedBeneficiary, AidClaim, BeneficiaryBackupDownload, OfflineSyncMetadata
from households.models import Household, Family, FamilyMember, Zone
from programs.models import Program, AidCategory, Assistance, EligibilityRule
from accounts.models import Barangay


def export_beneficiary_list(schedule):
    """
    Export a beneficiary list to .xlsx for offline backup.
    
    Args:
        schedule: AidSchedule instance with a beneficiary_list
        
    Returns:
        tuple: (filename, xlsx_bytes)
    """
    if not hasattr(schedule, 'beneficiary_list'):
        raise ValueError("Schedule has no beneficiary list")
    
    ben_list = schedule.beneficiary_list
    entries = ben_list.entries.select_related(
        'family', 'family__household', 'family__household__barangay', 
        'family__household__zone', 'family_member', 'household'
    ).all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Beneficiaries"
    
    # Metadata header
    metadata_data = [
        ["Schedule ID", schedule.id],
        ["Program", schedule.assistance.program.name if schedule.assistance else "N/A"],
        ["Aid Category", schedule.assistance.aid_category.name if schedule.assistance else "N/A"],
        ["Date/Time", schedule.schedule_datetime.strftime("%Y-%m-%d %H:%M")],
        ["Location", schedule.location],
        ["Prioritization Strategy", ben_list.prioritization_strategy_used],
        ["Budget", str(schedule.budget)],
        ["Per Beneficiary Amount", str(schedule.per_beneficiary_amount)],
        ["Generated At", ben_list.generated_at.strftime("%Y-%m-%d %H:%M")],
        ["", ""],  # Empty row separator
        ["Family ID", "Member ID", "Full Name", "RFID UID", "Address", "Eligibility Criteria Met"]
    ]
    
    # Write metadata
    for row_idx, row_data in enumerate(metadata_data, start=1):
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if row_idx <= 10:  # Metadata rows
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="E7F3FF", end_color="E7F3FF", fill_type="solid")
    
    # Write beneficiary entries
    row_idx = 12
    for entry in entries:
        if entry.family:
            family_id = entry.family.id
            member_id = entry.family_member.id if entry.family_member else ""
            rfid_uid = entry.family.rfid_uid or ""
            
            if entry.family_member:
                full_name = f"{entry.family_member.first_name} {entry.family_member.last_name}"
            else:
                full_name = f"{entry.family.family_name} Family"
            
            address = entry.family.household.address
        elif entry.household:
            family_id = ""
            member_id = ""
            rfid_uid = ""
            full_name = entry.household.address
            address = entry.household.address
        else:
            family_id = ""
            member_id = ""
            rfid_uid = ""
            full_name = "Unknown"
            address = ""
        
        # Determine eligibility criteria (simplified - could be expanded)
        criteria = "Listed beneficiary"
        if entry.added_manually:
            criteria += " (Manual Override)"
        
        ws.cell(row=row_idx, column=1, value=family_id)
        ws.cell(row=row_idx, column=2, value=member_id)
        ws.cell(row=row_idx, column=3, value=full_name)
        ws.cell(row=row_idx, column=4, value=rfid_uid)
        ws.cell(row=row_idx, column=5, value=address)
        ws.cell(row=row_idx, column=6, value=criteria)
        row_idx += 1
    
    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width
    
    # Generate filename
    date_str = schedule.schedule_datetime.strftime("%m-%d-%Y")
    program_name = schedule.assistance.program.name if schedule.assistance else "Unknown"
    aid_category = schedule.assistance.aid_category.name if schedule.assistance else "Unknown"
    filename = f"{date_str}_{program_name}_{aid_category}_Sched{schedule.id}.xlsx"
    
    # Save to bytes
    xlsx_bytes = BytesIO()
    wb.save(xlsx_bytes)
    xlsx_bytes.seek(0)
    
    return filename, xlsx_bytes.getvalue()


def export_full_offline_sync():
    """
    Export full municipal-wide data for offline sync.
    
    Returns a zip file containing:
    - households.json (Household, Family, FamilyMember, Zone data)
    - programs.json (Program, AidCategory, Assistance, EligibilityRule data)
    - claims.json (AidClaim history for cooldown/rotation eligibility)
    - metadata.json (sync timestamp, counts)
    
    Returns:
        tuple: (filename, zip_bytes)
    """
    zip_buffer = BytesIO()
    
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zipf:
        sync_timestamp = timezone.now().isoformat()
        
        # Export Households, Families, Members, Zones
        households_data = []
        for household in Household.objects.select_related('barangay', 'zone').prefetch_related('families__members').all():
            household_dict = {
                'id': household.id,
                'barangay_id': household.barangay.id,
                'barangay_name': household.barangay.name,
                'zone_id': household.zone.id,
                'zone_name': household.zone.name,
                'house_number': household.house_number,
                'land_use': household.land_use,
                'hazard_exposure': household.hazard_exposure,
                'flood_depth': household.flood_depth,
                'flood_frequency': household.flood_frequency,
                'hazard_other_description': household.hazard_other_description,
                'accessibility': household.accessibility,
                'location_notes': household.location_notes,
                'latitude': str(household.latitude) if household.latitude else None,
                'longitude': str(household.longitude) if household.longitude else None,
                'families': []
            }
            
            for family in household.families.filter(is_active=True, is_archived=False):
                family_dict = {
                    'id': family.id,
                    'family_name': family.family_name,
                    'rfid_uid': family.rfid_uid,
                    'is_active': family.is_active,
                    'is_archived': family.is_archived,
                    'members': []
                }
                
                for member in family.members.filter(date_of_death__isnull=True):
                    member_dict = {
                        'id': member.id,
                        'first_name': member.first_name,
                        'middle_name': member.middle_name,
                        'last_name': member.last_name,
                        'suffix': member.suffix,
                        'relationship': member.relationship,
                        'birthdate': member.birthdate.isoformat() if member.birthdate else None,
                        'sex': member.sex,
                        'civil_status': member.civil_status,
                        'monthly_income': str(member.monthly_income),
                        'is_pwd': member.is_pwd,
                        'is_solo_parent': member.is_solo_parent,
                        'is_senior_citizen': member.is_senior_citizen,
                        'is_indigenous': member.is_indigenous,
                        'is_out_of_school_youth': member.is_out_of_school_youth,
                        'is_out_of_school_children': member.is_out_of_school_children,
                    }
                    family_dict['members'].append(member_dict)
                
                household_dict['families'].append(family_dict)
            
            households_data.append(household_dict)
        
        zipf.writestr('households.json', json.dumps(households_data, indent=2))
        
        # Export Programs, AidCategories, Assistances, EligibilityRules
        programs_data = []
        for program in Program.objects.filter(is_active=True).prefetch_related('aid_categories__assistances__eligibility_rules').all():
            program_dict = {
                'id': program.id,
                'name': program.name,
                'description': program.description,
                'is_active': program.is_active,
                'is_emergency_program': program.is_emergency_program,
                'aid_categories': []
            }
            
            for aid_category in program.aid_categories.filter(is_active=True):
                category_dict = {
                    'id': aid_category.id,
                    'name': aid_category.name,
                    'description': aid_category.description,
                    'is_active': aid_category.is_active,
                    'assistances': []
                }
                
                for assistance in aid_category.assistances.filter(is_active=True):
                    assistance_dict = {
                        'id': assistance.id,
                        'beneficiary_type': assistance.beneficiary_type,
                        'aid_type': assistance.aid_type,
                        'minimum_age': assistance.minimum_age,
                        'requires_pwd': assistance.requires_pwd,
                        'requires_solo_parent': assistance.requires_solo_parent,
                        'requires_senior_citizen': assistance.requires_senior_citizen,
                        'is_active': assistance.is_active,
                        'eligibility_rules': []
                    }
                    
                    for rule in assistance.eligibility_rules.filter(is_active=True):
                        rule_dict = {
                            'id': rule.id,
                            'rule_type': rule.rule_type,
                            'config': rule.config,
                            'is_active': rule.is_active
                        }
                        assistance_dict['eligibility_rules'].append(rule_dict)
                    
                    category_dict['assistances'].append(assistance_dict)
                
                program_dict['aid_categories'].append(category_dict)
            
            programs_data.append(program_dict)
        
        zipf.writestr('programs.json', json.dumps(programs_data, indent=2))
        
        # Export Claims history (for cooldown/rotation eligibility)
        claims_data = []
        for claim in AidClaim.objects.select_related(
            'family', 'family_member', 'assistance__program', 'schedule'
        ).all():
            claim_dict = {
                'id': claim.id,
                'family_id': claim.family.id,
                'family_member_id': claim.family_member.id if claim.family_member else None,
                'assistance_id': claim.assistance.id if claim.assistance else None,
                'program_name': claim.assistance.program.name if claim.assistance	else None,
                'schedule_id': claim.schedule.id if claim.schedule else None,
                'claim_type': claim.claim_type,
                'claimed_at': claim.claimed_at.isoformat(),
                'amount': str(claim.amount) if claim.amount else None,
                'is_late_scheduled_claim': claim.is_late_scheduled_claim,
            }
            claims_data.append(claim_dict)
        
        zipf.writestr('claims.json', json.dumps(claims_data, indent=2))
        
        # Export typhoon signal status for offline eligibility checks
        from households.models import WeatherSnapshot
        latest_snapshot = WeatherSnapshot.objects.order_by('-timestamp').first()
        typhoon_data = {
            'has_snapshot': latest_snapshot is not None,
            'timestamp': latest_snapshot.timestamp.isoformat() if latest_snapshot else None,
            'wind_speed_kph': latest_snapshot.wind_speed_kph if latest_snapshot else None,
            'signal_level': latest_snapshot.signal_level if latest_snapshot else None,
        }
        zipf.writestr('typhoon_status.json', json.dumps(typhoon_data, indent=2))
        
        # Export metadata
        metadata = {
            'sync_timestamp': sync_timestamp,
            'households_count': len(households_data),
            'programs_count': len(programs_data),
            'claims_count': len(claims_data),
            'includes_typhoon_status': True,
        }
        zipf.writestr('metadata.json', json.dumps(metadata, indent=2))
    
    zip_buffer.seek(0)
    filename = f"offline_sync_{timezone.now().strftime('%Y%m%d_%H%M%S')}.zip"
    
    return filename, zip_buffer.getvalue()


def import_beneficiary_list(xlsx_file, user):
    """
    Import a beneficiary list from .xlsx on the offline instance.
    Creates the AidSchedule and beneficiary list/entries.
    
    Args:
        xlsx_file: Uploaded .xlsx file
        user: User performing the import (for audit logging)
        
    Returns:
        AidSchedule instance that was created
    """
    from openpyxl import load_workbook
    
    wb = load_workbook(xlsx_file)
    ws = wb.active
    
    # Read metadata (rows 1-10)
    metadata = {}
    for row_idx in range(1, 11):
        cell = ws.cell(row=row_idx, column=1)
        value_cell = ws.cell(row=row_idx, column=2)
        if cell.value and value_cell.value:
            metadata[cell.value] = value_cell.value
    
    # Parse schedule data
    program_name = metadata.get("Program", "")
    aid_category_name = metadata.get("Aid Category", "")
    schedule_datetime_str = metadata.get("Date/Time", "")
    location = metadata.get("Location", "")
    prioritization_strategy = metadata.get("Prioritization Strategy", "LOWEST_INCOME_FIRST")
    budget = Decimal(metadata.get("Budget", "0"))
    per_beneficiary_amount = Decimal(metadata.get("Per Beneficiary Amount", "0"))
    
    # Find or create Program, AidCategory, Assistance
    program, _ = Program.objects.get_or_create(
        name=program_name,
        defaults={'description': 'Imported from offline backup', 'is_active': True}
    )
    
    aid_category, _ = AidCategory.objects.get_or_create(
        program=program,
        name=aid_category_name,
        defaults={'description': 'Imported from offline backup', 'is_active': True}
    )
    
    assistance, _ = Assistance.objects.get_or_create(
        program=program,
        aid_category=aid_category,
        beneficiary_type='family',  # Default to family for simplicity
        defaults={'is_active': True}
    )
    
    # Parse datetime
    from datetime import datetime
    schedule_datetime = datetime.strptime(schedule_datetime_str, "%Y-%m-%d %H:%M")
    if timezone.is_naive(schedule_datetime):
        schedule_datetime = timezone.make_aware(schedule_datetime)
    
    # Create AidSchedule
    schedule = AidSchedule.objects.create(
        assistance=assistance,
        schedule_datetime=schedule_datetime,
        location=location,
        budget=budget,
        per_beneficiary_amount=per_beneficiary_amount,
        prioritization_strategy=prioritization_strategy,
        created_by=user
    )
    
    # Create beneficiary list
    ben_list = GeneratedBeneficiaryList.objects.create(
        schedule=schedule,
        generated_by=user,
        prioritization_strategy_used=prioritization_strategy
    )
    
    # Read beneficiary entries (starting from row 12)
    row_idx = 12
    while True:
        family_id = ws.cell(row=row_idx, column=1).value
        if not family_id:
            break
        
        member_id = ws.cell(row=row_idx, column=2).value
        full_name = ws.cell(row=row_idx, column=3).value
        rfid_uid = ws.cell(row=row_idx, column=4).value
        address = ws.cell(row=row_idx, column=5).value
        criteria = ws.cell(row=row_idx, column=6).value
        
        # Try to find family by ID or RFID
        family = None
        family_member = None
        
        if family_id:
            try:
                family = Family.objects.get(id=family_id)
            except Family.DoesNotExist:
                pass
        
        if not family and rfid_uid:
            try:
                family = Family.objects.get(rfid_uid=rfid_uid)
            except Family.DoesNotExist:
                pass
        
        if family and member_id:
            try:
                family_member = FamilyMember.objects.get(id=member_id, family=family)
            except FamilyMember.DoesNotExist:
                pass
        
        # Create beneficiary entry
        is_manual = "Manual Override" in str(criteria) if criteria else False
        GeneratedBeneficiary.objects.create(
            beneficiary_list=ben_list,
            family=family,
            family_member=family_member,
            added_manually=is_manual,
            added_by=user if is_manual else None
        )
        
        row_idx += 1
    
    return schedule


def import_full_offline_sync(zip_file, user):
    """
    Import full municipal-wide data from zip on offline instance.
    
    Args:
        zip_file: Uploaded .zip file
        user: User performing the import
        
    Returns:
        dict with import statistics
    """
    with zipfile.ZipFile(zip_file, 'r') as zipf:
        # Read metadata
        metadata = json.loads(zipf.read('metadata.json').decode('utf-8'))
        
        # Import typhoon signal status if available
        if 'typhoon_status.json' in zipf.namelist():
            typhoon_data = json.loads(zipf.read('typhoon_status.json').decode('utf-8'))
            if typhoon_data.get('has_snapshot') and typhoon_data.get('timestamp'):
                from households.models import WeatherSnapshot
                from datetime import datetime
                WeatherSnapshot.objects.create(
                    timestamp=datetime.fromisoformat(typhoon_data['timestamp']),
                    wind_speed_kph=typhoon_data.get('wind_speed_kph'),
                    signal_level=typhoon_data.get('signal_level')
                )
        
        # Import households
        households_data = json.loads(zipf.read('households.json').decode('utf-8'))
        households_imported = 0
        families_imported = 0
        members_imported = 0
        
        for household_data in households_data:
            # Get or create barangay and zone
            barangay, _ = Barangay.objects.get_or_create(
                name=household_data['barangay_name'],
                defaults={}
            )
            
            zone, _ = Zone.objects.get_or_create(
                barangay=barangay,
                name=household_data['zone_name'],
                defaults={}
            )
            
            # Get or create household
            household, created = Household.objects.get_or_create(
                id=household_data['id'],
                defaults={
                    'barangay': barangay,
                    'zone': zone,
                    'house_number': household_data['house_number'],
                    'land_use': household_data['land_use'],
                    'hazard_exposure': household_data['hazard_exposure'],
                    'flood_depth': household_data['flood_depth'],
                    'flood_frequency': household_data['flood_frequency'],
                    'hazard_other_description': household_data['hazard_other_description'],
                    'accessibility': household_data['accessibility'],
                    'location_notes': household_data['location_notes'],
                    'latitude': Decimal(household_data['latitude']) if household_data['latitude'] else None,
                    'longitude': Decimal(household_data['longitude']) if household_data['longitude'] else None,
                }
            )
            if created:
                households_imported += 1
            
            # Import families
            for family_data in household_data['families']:
                family, created = Family.objects.get_or_create(
                    id=family_data['id'],
                    defaults={
                        'household': household,
                        'family_name': family_data['family_name'],
                        'rfid_uid': family_data['rfid_uid'],
                        'is_active': family_data['is_active'],
                        'is_archived': family_data['is_archived'],
                    }
                )
                if created:
                    families_imported += 1
                
                # Import members
                for member_data in family_data['members']:
                    member, created = FamilyMember.objects.get_or_create(
                        id=member_data['id'],
                        defaults={
                            'family': family,
                            'first_name': member_data['first_name'],
                            'middle_name': member_data['middle_name'],
                            'last_name': member_data['last_name'],
                            'suffix': member_data['suffix'],
                            'relationship': member_data['relationship'],
                            'birthdate': datetime.fromisoformat(member_data['birthdate']).date() if member_data['birthdate'] else None,
                            'sex': member_data['sex'],
                            'civil_status': member_data['civil_status'],
                            'monthly_income': Decimal(member_data['monthly_income']),
                            'is_pwd': member_data['is_pwd'],
                            'is_solo_parent': member_data['is_solo_parent'],
                            'is_senior_citizen': member_data['is_senior_citizen'],
                            'is_indigenous': member_data['is_indigenous'],
                            'is_out_of_school_youth': member_data['is_out_of_school_youth'],
                            'is_out_of_school_children': member_data['is_out_of_school_children'],
                        }
                    )
                    if created:
                        members_imported += 1
        
        # Import programs
        programs_data = json.loads(zipf.read('programs.json').decode('utf-8'))
        programs_imported = 0
        categories_imported = 0
        assistances_imported = 0
        
        for program_data in programs_data:
            program, created = Program.objects.get_or_create(
                id=program_data['id'],
                defaults={
                    'name': program_data['name'],
                    'description': program_data['description'],
                    'is_active': program_data['is_active'],
                    'is_emergency_program': program_data['is_emergency_program'],
                }
            )
            if created:
                programs_imported += 1
            
            for category_data in program_data['aid_categories']:
                category, created = AidCategory.objects.get_or_create(
                    id=category_data['id'],
                    defaults={
                        'program': program,
                        'name': category_data['name'],
                        'description': category_data['description'],
                        'is_active': category_data['is_active'],
                    }
                )
                if created:
                    categories_imported += 1
                
                for assistance_data in category_data['assistances']:
                    assistance, created = Assistance.objects.get_or_create(
                        id=assistance_data['id'],
                        defaults={
                            'program': program,
                            'aid_category': category,
                            'beneficiary_type': assistance_data['beneficiary_type'],
                            'aid_type': assistance_data['aid_type'],
                            'minimum_age': assistance_data['minimum_age'],
                            'requires_pwd': assistance_data['requires_pwd'],
                            'requires_solo_parent': assistance_data['requires_solo_parent'],
                            'requires_senior_citizen': assistance_data['requires_senior_citizen'],
                            'is_active': assistance_data['is_active'],
                        }
                    )
                    if created:
                        assistances_imported += 1
                    
                    # Import eligibility rules
                    for rule_data in assistance_data['eligibility_rules']:
                        EligibilityRule.objects.get_or_create(
                            id=rule_data['id'],
                            defaults={
                                'assistance': assistance,
                                'rule_type': rule_data['rule_type'],
                                'config': rule_data['config'],
                                'is_active': rule_data['is_active'],
                            }
                        )
        
        # Import claims
        claims_data = json.loads(zipf.read('claims.json').decode('utf-8'))
        claims_imported = 0
        
        for claim_data in claims_data:
            try:
                family = Family.objects.get(id=claim_data['family_id'])
                family_member = FamilyMember.objects.filter(id=claim_data['family_member_id']).first() if claim_data['family_member_id'] else None
                assistance = Assistance.objects.filter(id=claim_data['assistance_id']).first() if claim_data['assistance_id'] else None
                schedule = AidSchedule.objects.filter(id=claim_data['schedule_id']).first() if claim_data['schedule_id'] else None
                
                claim, created = AidClaim.objects.get_or_create(
                    id=claim_data['id'],
                    defaults={
                        'family': family,
                        'family_member': family_member,
                        'assistance': assistance,
                        'schedule': schedule,
                        'claim_type': claim_data['claim_type'],
                        'claimed_at': datetime.fromisoformat(claim_data['claimed_at']),
                        'amount': Decimal(claim_data['amount']) if claim_data['amount'] else None,
                        'is_late_scheduled_claim': claim_data['is_late_scheduled_claim'],
                    }
                )
                if created:
                    claims_imported += 1
            except (Family.DoesNotExist, Assistance.DoesNotExist, AidSchedule.DoesNotExist):
                # Skip claims with missing references
                continue
        
        # Create or update OfflineSyncMetadata record
        sync_metadata, _ = OfflineSyncMetadata.objects.get_or_create(
            synced_by=user,
            defaults={
                'synced_at': timezone.now(),
                'households_count': households_imported,
                'programs_count': programs_imported,
                'claims_count': claims_imported,
            }
        )
        # Update if record already exists
        if not _:
            sync_metadata.synced_at = timezone.now()
            sync_metadata.households_count = households_imported
            sync_metadata.programs_count = programs_imported
            sync_metadata.claims_count = claims_imported
            sync_metadata.save()
        
        return {
            'sync_timestamp': metadata['sync_timestamp'],
            'households_imported': households_imported,
            'families_imported': families_imported,
            'members_imported': members_imported,
            'programs_imported': programs_imported,
            'categories_imported': categories_imported,
            'assistances_imported': assistances_imported,
            'claims_imported': claims_imported,
        }


def export_offline_claims(schedule_id=None, walkin_only=False):
    """
    Export claims from offline instance for reconciliation.
    
    Args:
        schedule_id: If provided, export only claims for this schedule
        walkin_only: If True, export only walk-in claims
        
    Returns:
        tuple: (filename, json_bytes)
    """
    claims = AidClaim.objects.select_related(
        'family', 'family_member', 'assistance__program', 'assistance__aid_category', 'schedule'
    )
    
    if schedule_id:
        claims = claims.filter(schedule_id=schedule_id)
    
    if walkin_only:
        claims = claims.filter(claim_type='WALK_IN')
    
    # Only export unexported claims
    claims = claims.filter(is_exported=False)
    
    claims_data = []
    for claim in claims:
        claim_dict = {
            'offline_claim_id': claim.id,
            'family_id': claim.family.id,
            'family_member_id': claim.family_member.id if claim.family_member else None,
            'assistance_id': claim.assistance.id if claim.assistance else None,
            'program_name': claim.assistance.program.name if claim.assistance else None,
            'aid_category_name': claim.assistance.aid_category.name if claim.assistance else None,
            'schedule_id': claim.schedule.id if claim.schedule else None,
            'claim_type': claim.claim_type,
            'claimed_at': claim.claimed_at.isoformat(),
            'amount': str(claim.amount) if claim.amount else None,
            'is_late_scheduled_claim': claim.is_late_scheduled_claim,
            'original_schedule_id': claim.original_schedule.id if claim.original_schedule else None,
        }
        claims_data.append(claim_dict)
    
    # Mark as exported
    claims.update(is_exported=True)
    
    filename = f"offline_claims_export_{timezone.now().strftime('%Y%m%d_%H%M%S')}.json"
    json_bytes = json.dumps(claims_data, indent=2).encode('utf-8')
    
    return filename, json_bytes
